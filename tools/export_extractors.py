"""把当前快照里的全部提取器导出成 Excel，用于前后对比。

用法：
    python tools/export_extractors.py                # 用最近一次总览快照
    python tools/export_extractors.py --out 目录      # 指定输出目录（默认 exports/）

注意：读的是程序里的总览快照（data/pi.db 的 overview_cache），不是实时 ESI。
导出前先在页面上点一次「刷新」，脚本会把快照时间写进文件名和「说明」页。
"""

from __future__ import annotations

import argparse
import json
import shutil
import sqlite3
import sys
import tempfile
import time
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import paths

# ---------------------------------------------------------------- 样式
HDR_FILL = PatternFill("solid", fgColor="1C2938")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONO = Font(name="Consolas", size=10)


def style_header(ws, ncol: int, height: int = 26) -> None:
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = height
    ws.freeze_panes = "A2"


def load_snapshot() -> tuple[float, dict]:
    """读最近一次总览快照。复制一份库再读，避免和正在运行的程序抢锁。"""
    db_path = paths.data_path("data", "pi.db")
    if not db_path.exists():
        raise SystemExit(f"找不到数据库：{db_path}\n先启动程序并刷新一次。")
    with tempfile.TemporaryDirectory() as tmp:
        tmp_db = Path(tmp) / "pi.db"
        shutil.copy(db_path, tmp_db)
        conn = sqlite3.connect(tmp_db)
        conn.row_factory = sqlite3.Row
        row = conn.execute("select ts, data from overview_cache order by ts desc limit 1").fetchone()
        conn.close()
    if not row:
        raise SystemExit("还没有任何总览快照。先在页面上点一次「刷新」。")
    return row["ts"], json.loads(row["data"])


DETAIL_HEADERS = [
    ("角色", 12), ("星系", 11), ("星球", 14), ("星球类型", 10), ("安全", 7),
    ("产物", 11), ("产物(英文)", 18), ("层级", 6),
    ("采矿头", 7), ("周期(分)", 9), ("每周期产量", 12), ("每小时产量", 12),
    ("程序开始", 17), ("程序结束", 17), ("程序时长(h)", 11),
    ("已运行(h)", 10), ("进度", 8), ("剩余(h)", 9),
    ("全程外推", 12), ("剩余可采", 12),
    ("仓储占用", 9), ("数据时间", 17),
    ("角色ID", 12), ("引脚ID", 15),
]


def build_detail(ws, ov: dict) -> int:
    for i, (name, width) in enumerate(DETAIL_HEADERS, start=1):
        ws.cell(row=1, column=i, value=name)
        ws.column_dimensions[get_column_letter(i)].width = width
    style_header(ws, len(DETAIL_HEADERS))

    r = 2
    for c in ov["colonies"]:
        for e in c["extractors"]:
            expiry = e.get("expiry")
            dur = e.get("duration") or 0
            install = (expiry - dur) if (expiry and dur) else None
            vals = [
                c["character_name"], c["system_name"], c["planet_name"],
                c["planet_type"], c["security"],
                e["product"], e.get("product_en", ""), e.get("tier", ""),
                e.get("heads"), round((e.get("cycle_time") or 0) / 60),
                e.get("qty_per_cycle"), round(e.get("rate_per_hour") or 0),
                datetime.fromtimestamp(install) if install else None,
                datetime.fromtimestamp(expiry) if expiry else None,
                round(dur / 3600, 1) if dur else None,
                round((dur - (e.get("remaining") or 0)) / 3600, 1) if dur else None,
                e.get("progress"),
                round((e.get("remaining") or 0) / 3600, 1),
                round(e.get("est_total") or 0),
                round(e.get("est_remaining") or 0),
                c.get("fill"),
                datetime.fromtimestamp(c["data_time"]) if c.get("data_time") else None,
                c["character_id"], e["pin_id"],
            ]
            for i, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=i, value=v)
                cell.border = BORDER
                cell.font = Font(size=10)
                if i in (13, 14, 22):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                elif i == 17:
                    cell.number_format = "0.0%"
                elif i == 21:
                    cell.number_format = "0%"
                elif i in (11, 12, 19, 20):
                    cell.number_format = "#,##0"
                elif i in (15, 16, 18):
                    cell.number_format = "0.0"
                elif i in (23, 24):
                    cell.font = MONO

            # 剩余时间上色：<2h 橙、已到期 红
            rem_h = (e.get("remaining") or 0) / 3600
            rem_cell = ws.cell(row=r, column=18)
            if rem_h <= 0:
                rem_cell.font = Font(size=10, color="C00000", bold=True)
                ws.cell(row=r, column=20).font = Font(size=10, color="C00000")
            elif rem_h < 2:
                rem_cell.font = Font(size=10, color="BF8F00", bold=True)
            r += 1

    ws.auto_filter.ref = f"A1:{get_column_letter(len(DETAIL_HEADERS))}{r - 1}"
    ws.sheet_view.zoomScale = 90
    return r - 2


def build_summary(ws, ov: dict) -> None:
    headers = [("产物", 12), ("层级", 7), ("提取器数", 10), ("采矿头合计", 11),
               ("每小时合计", 14), ("剩余可采合计", 15), ("全程外推合计", 15)]
    for i, (name, width) in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=name)
        ws.column_dimensions[get_column_letter(i)].width = width
    style_header(ws, len(headers))

    agg: dict[str, dict] = {}
    for c in ov["colonies"]:
        for e in c["extractors"]:
            a = agg.setdefault(e["product"], {"tier": e.get("tier"), "n": 0, "heads": 0,
                                              "rate": 0.0, "rem": 0.0, "total": 0.0})
            a["n"] += 1
            a["heads"] += e.get("heads") or 0
            a["rate"] += e.get("rate_per_hour") or 0
            a["rem"] += e.get("est_remaining") or 0
            a["total"] += e.get("est_total") or 0

    r = 2
    for name, a in sorted(agg.items(), key=lambda kv: -kv[1]["rate"]):
        for i, v in enumerate([name, a["tier"], a["n"], a["heads"],
                               round(a["rate"]), round(a["rem"]), round(a["total"])], start=1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.border = BORDER
            cell.font = Font(size=10)
            if i >= 5:
                cell.number_format = "#,##0"
        r += 1

    for col, letter in ((1, "A"), (3, "C"), (4, "D"), (5, "E")):
        cell = ws.cell(row=r, column=col, value="合计" if col == 1 else f"=SUM({letter}2:{letter}{r - 1})")
        cell.font = Font(size=10, bold=True)
        if col == 5:
            cell.number_format = "#,##0"
    for i in range(1, 8):
        ws.cell(row=r, column=i).border = BORDER
        ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor="EEF2F7")


def build_notes(ws, ov: dict, stamp: datetime) -> None:
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 96
    n_ext = sum(len(c["extractors"]) for c in ov["colonies"])
    n_char = len({c["character_name"] for c in ov["colonies"]})
    notes = [
        ("快照时间", stamp.strftime("%Y-%m-%d %H:%M:%S") + "（本机时间）"),
        ("数据来源", "ESI /characters/{id}/planets/ 与 /planets/{planet_id}/，经行星工业监视器解析后导出"),
        ("覆盖范围", f"{len(ov['colonies'])} 个殖民地、{n_ext} 个提取器、{n_char} 个角色"),
        ("", ""),
        ("每周期产量", "ESI 的 qty_per_cycle。注意：它是「最近一次模拟那个周期实际采到多少」，"
                       "不是全程平均；会随矿脉衰减一路变小。程序刚开时最接近峰值。"),
        ("全程外推", "= 每周期产量 × 总周期数。因为用的是末次速率，会低于游戏面板上的「总量」（实测某提取器低 17%）。"),
        ("剩余可采", "= 每周期产量 × 剩余周期数。偏乐观，实际会因继续衰减而略低。"),
        ("进度 / 已运行", "按程序开始时间到快照时间的比例计算，与游戏里采集面板的进度条一致。"),
        ("数据时间", "殖民地列表里的 last_update，即 CCP 最后一次重算该殖民地的时间。"
                     "比快照时间早多少，说明那块殖民地的数据有多旧。"),
        ("比对方法", "下次导出后用「角色ID + 引脚ID」做主键对比（同一程序内引脚ID不变）。"
                     "若程序已重开，引脚ID会变，需按「角色+星球+产物」匹配。"),
        ("", ""),
        ("已知问题", "qty_per_cycle 只有在有人（客户端）打开过该殖民地后才更新，"
                     "长时间不看的殖民地会停在旧值上。做对比前建议先在游戏里过一遍。"),
    ]
    ws.cell(row=1, column=1, value="提取器快照说明").font = Font(size=13, bold=True)
    for i, (k, v) in enumerate(notes, start=3):
        a = ws.cell(row=i, column=1, value=k)
        a.font = Font(size=10, bold=True)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=i, column=2, value=v)
        b.font = Font(size=10)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if v:
            ws.row_dimensions[i].height = max(15, 15 * (len(v) // 60 + 1))


def main() -> None:
    ap = argparse.ArgumentParser(description="导出提取器快照到 Excel")
    ap.add_argument("--out", default="exports", help="输出目录（默认 exports/）")
    args = ap.parse_args()

    snapshot_ts, ov = load_snapshot()
    stamp = datetime.fromtimestamp(snapshot_ts)
    age_min = (time.time() - snapshot_ts) / 60

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"提取器快照_{stamp:%Y-%m-%d_%H%M}.xlsx"

    wb = Workbook()
    n = build_detail(wb.active, ov)
    wb.active.title = "提取器明细"
    build_summary(wb.create_sheet("按产物汇总"), ov)
    build_notes(wb.create_sheet("说明"), ov, stamp)
    wb.save(out_path)

    print(f"快照时间：{stamp:%Y-%m-%d %H:%M}（{age_min:.0f} 分钟前）")
    if age_min > 60:
        print("  提醒：快照超过 1 小时了，先在页面上点「刷新」再导出，数据才是当前的。")
    print(f"导出 {n} 个提取器 → {out_path}")


if __name__ == "__main__":
    main()
